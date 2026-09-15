#!/usr/bin/env python3
"""Normalize a heterogeneous SRT reference tape into the canonical registry schema.

L1 execution adapter. Standard-library-first (CSV native; XLSX only if openpyxl is
installed). Mapping-driven, so a new issuer is a new JSON mapping, not new code.

Refuses to fabricate (mirrors the skill's hard stops):
  - required field missing on an obligor row     -> STOP (exit 1)
  - a credit value is impossible (pd/recovery/lgd outside [0,1]; rona < 0) -> STOP
  - a mapped column is missing from the tape      -> STOP if required, else WARN
  - SUM(rona) does not tie to the stated notional -> STOP (exit 1) WHEN a stated notional is
    configured (`deal_meta.stated_reference_notional` > 0); with the template default of 0
    the tie-out is SKIPPED and reported as such, not stopped. Tolerance: `tie_out.tolerance_pct`
    (default 0.5 %).
Which fields are required is set by the mapping (`required_fields`; template default
["rona", "maturity_date"]). PD is required only when the mapping lists it -- a registry
that carries LGD but no PD (the H11 / defaulted-tab path) is a documented use.
Recovery->LGD is the only credit derivation; PD/recovery are never invented (missing
values are reported for the stratification fallback under recorded analyst judgment).
Non-obligor rows (totals/footers with no obligor identity) are skipped, not ingested.
A status value outside `status_rules.defaulted_values` is treated as Performing and the
distinct values seen are listed in the report notes.
On a STOP the canonical rows are written to `<out>.STOPPED.csv`, never to `<out>`, so a
downstream step that checks for the file cannot proceed on a stopped run.

Usage:
  python normalize_tape.py --tape RAW.csv|RAW.xlsx --mapping issuer-mapping.json \
      --out canonical-registry.csv [--report normalize-report.json]

See adapters/canonical-registry-schema.md and adapters/issuer-mapping.template.json.
"""
import argparse
import csv
import json
import os
import sys

CANON_COLS = ["obligor_id", "obligor_name", "country_domicile", "country_incorporation",
              "industry", "rating", "rating_bucket", "maturity_date", "currency", "rona",
              "pd", "recovery", "lgd", "seniority", "secured", "status", "source_row"]

# Built-in rating buckets (S&P + Moody's, UPPER). Merged with the mapping's map so an
# unanticipated scale is not silently mis-bucketed (e.g. Moody's 'Baa2' -> IG, not 'B').
DEFAULT_BUCKET_MAP = {
    "AAA": "IG", "AA": "IG", "A": "IG", "BBB": "IG", "BAA": "IG",
    "BB": "BB", "BA": "BB", "B": "B",
    "CCC": "CCC", "CC": "CCC", "C": "CCC", "CAA": "CCC", "CA": "CCC", "D": "CCC",
    "NR": "NA", "WR": "NA", "": "NA",
}

# Obligor-identity values that mark a non-obligor row (totals/footers). Matched
# case-insensitively; overridable/extendable via the mapping's `non_obligor_labels`.
DEFAULT_NONOBLIGOR = {"", "total", "totals", "grand total", "subtotal", "sub-total",
                      "aggregate", "portfolio total", "sum"}


def read_rows(tape_path, source):
    """Return (rows, headers) from a CSV or XLSX tape."""
    stype = (source.get("type") or "").lower()
    if stype == "csv" or tape_path.lower().endswith(".csv"):
        with open(tape_path, newline="", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            rows = list(reader)
            return rows, [h for h in (reader.fieldnames or [])]
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("ERROR: tape is xlsx but openpyxl is not installed. "
                 "Install openpyxl, or export the tab to CSV and set source.type=csv.")
    wb = load_workbook(tape_path, read_only=True, data_only=True)
    sheet = source.get("registry_sheet")
    if sheet and sheet not in wb.sheetnames:
        sys.exit(f"ERROR: registry_sheet '{sheet}' not found. Tabs present: {wb.sheetnames}")
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    header_row = int(source.get("header_row", 1))
    rows, headers = [], None
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i < header_row:
            continue
        if i == header_row:
            headers = [str(c).strip() if c is not None else "" for c in row]
            continue
        if all(c is None for c in row):
            continue
        # ragged rows: pad short rows with None instead of IndexError
        rows.append({headers[j]: (row[j] if j < len(row) else None) for j in range(len(headers))})
    return rows, (headers or [])


def to_float(v, scale=1.0):
    """Parse a number. A trailing '%' is treated as percent (divided by 100)."""
    if v is None:
        return None
    s0 = str(v).strip()
    if s0 == "":
        return None
    pct = s0.endswith("%")
    s = s0.replace(",", "").replace("%", "")
    try:
        x = float(s)
    except ValueError:
        return None
    if pct:
        x /= 100.0
    return x * scale


def to_iso(v, fmts):
    if v is None:
        return None
    s = str(v).strip()
    if s == "":
        return None
    if len(s) >= 10 and s[4] == "-" and s[7] == "-":  # already ISO-ish (incl. openpyxl datetime str)
        return s[:10]
    from datetime import datetime
    for fmt in fmts:
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None  # unparseable -> flagged as missing downstream


def rating_bucket(raw, bmap):
    if raw is None:
        return "NA"
    s = str(raw).strip().upper()
    if s == "":
        return "NA"
    for token in sorted([t for t in bmap if t], key=len, reverse=True):  # longest token first
        if s.startswith(token):
            return bmap[token]
    return bmap.get("", "NA")


def load_mapping(mapping_path):
    with open(mapping_path, encoding="utf-8") as f:
        return json.load(f)


def mapping_context(mapping):
    transforms = mapping.get("transforms", {})
    return {
        "cmap": {
            k: v
            for k, v in mapping.get("column_map", {}).items()
            if v and not k.startswith("_")
        },
        "transforms": transforms,
        "date_formats": transforms.get("date_formats", ["%Y-%m-%d", "%d-%b-%y", "%m/%d/%Y"]),
        "bucket_map": {**DEFAULT_BUCKET_MAP, **transforms.get("rating_bucket_map", {})},
        "defaulted": {
            s.lower()
            for s in mapping.get("status_rules", {}).get("defaulted_values", ["defaulted"])
        },
        "required": mapping.get("required_fields", ["rona", "maturity_date"]),
        "meta": mapping.get("deal_meta", {}),
        "tie": mapping.get("tie_out", {}),
        "non_obligor": DEFAULT_NONOBLIGOR
        | {s.strip().lower() for s in mapping.get("non_obligor_labels", [])},
    }


def validate_mapped_columns(headers, cmap, required):
    stops, notes, skipped = [], [], 0
    hdr_set = set(headers)
    for field, src in cmap.items():
        if headers and src not in hdr_set:
            if field in required:
                stops.append(f"mapped column for required field '{field}' ('{src}') not found in tape header")
            else:
                notes.append(f"mapped column '{src}' (-> {field}) not found in tape header; '{field}' left blank")
    return stops, notes, skipped


def source_value(raw_row, cmap, field):
    return raw_row.get(cmap[field]) if field in cmap else None


def canonical_row(raw_row, idx, context, seen_ids):
    def get(field):
        return source_value(raw_row, context["cmap"], field)

    tf = context["transforms"]
    rec = to_float(get("recovery"), tf.get("recovery_scale", 1.0))
    lgd_src = to_float(get("lgd")) if "lgd" in context["cmap"] else None
    seniority = get("seniority")
    secured = None
    if seniority:
        sl = str(seniority).lower()
        secured = ("unsecured" not in sl) and ("secured" in sl)

    st_raw = get("status")
    status = (
        "Defaulted"
        if st_raw is not None and str(st_raw).strip().lower() in context["defaulted"]
        else "Performing"
    )
    ident = get("obligor_id") or get("obligor_name")
    raw_id = get("obligor_id")
    base_id = (
        str(raw_id).strip()
        if raw_id not in (None, "")
        else (str(ident).strip() if ident else f"row{idx}")
    )
    seen_ids[base_id] = seen_ids.get(base_id, 0) + 1
    obligor_id = base_id if seen_ids[base_id] == 1 else f"{base_id} #{seen_ids[base_id]}"

    return {
        "obligor_id": obligor_id,
        "obligor_name": get("obligor_name"),
        "country_domicile": get("country_domicile"),
        "country_incorporation": get("country_incorporation"),
        "industry": get("industry"),
        "rating": get("rating"),
        "rating_bucket": rating_bucket(get("rating"), context["bucket_map"]),
        "maturity_date": to_iso(get("maturity_date"), context["date_formats"]),
        "currency": context["meta"].get("currency"),
        "rona": to_float(get("rona"), tf.get("rona_scale", 1.0)),
        "pd": to_float(get("pd"), tf.get("pd_scale", 1.0)),
        "recovery": rec,
        "lgd": (1.0 - rec) if rec is not None else lgd_src,
        "seniority": seniority,
        "secured": secured,
        "status": status,
        "source_row": idx,
    }


def validate_canonical_row(row, required):
    stops, notes = [], []
    label = row["obligor_id"]
    for req in required:
        if row.get(req) is None:
            stops.append(f"{label}: missing required field '{req}'")
    for fld in ("pd", "recovery", "lgd"):
        val = row.get(fld)
        if val is not None and (val < -1e-9 or val > 1.0 + 1e-6):
            stops.append(f"{label}: {fld}={val:g} outside [0,1] (check percent vs decimal / source)")
    if row.get("rona") is not None and row["rona"] < 0:
        stops.append(f"{label}: rona={row['rona']:g} is negative")
    if row.get("recovery") is None and row.get("lgd") is None:
        notes.append(f"{label}: no recovery/LGD (stratification fallback under recorded judgment)")
    return stops, notes


def tie_out_summary(out_rows, context, stops, notes):
    meta, tie = context["meta"], context["tie"]
    count_def = tie.get("count_defaulted_in_total", True)
    summed = sum(
        r["rona"]
        for r in out_rows
        if r["rona"] is not None and (count_def or r["status"] != "Defaulted")
    )
    stated = meta.get("stated_reference_notional")
    tol = tie.get("tolerance_pct", 0.005)
    if stated is None or stated <= 0:
        notes.append(
            "tie-out skipped: no stated_reference_notional configured "
            "(set it to enable the RONA reconciliation)"
        )
        return summed, stated, "SKIPPED", f"SUM(rona)={summed:,.6g}; stated notional not configured -> tie-out SKIPPED"

    gap = abs(summed - stated) / stated
    status = "OK" if gap <= tol else "FAIL"
    message = f"SUM(rona)={summed:,.6g} vs stated {stated:,.6g} (gap {gap:.2%}, tol {tol:.2%}) -> {status}"
    if status == "FAIL":
        stops.append("RONA does not reconcile to the stated notional")
    return summed, stated, status, message


def normalize_rows(raw, headers, mapping):
    context = mapping_context(mapping)
    cmap = context["cmap"]
    stops, notes, skipped = validate_mapped_columns(headers, cmap, context["required"])
    identity_mapped = ("obligor_id" in cmap) or ("obligor_name" in cmap)
    seen_ids = {}
    out_rows = []
    for idx, r in enumerate(raw, start=2):
        ident = source_value(r, cmap, "obligor_id") or source_value(r, cmap, "obligor_name")
        ident_norm = str(ident).strip().lower() if ident is not None else ""
        if identity_mapped and ident_norm in context["non_obligor"]:
            skipped += 1  # totals/footer/blank/labelled-aggregate row -> not an obligor
            continue

        row = canonical_row(r, idx, context, seen_ids)
        out_rows.append(row)
        row_stops, row_notes = validate_canonical_row(row, context["required"])
        stops.extend(row_stops)
        notes.extend(row_notes)

    # review 2026-09-11 (C11): a repeated obligor id is renamed "<id> #2"; say so.
    for base_id, n in seen_ids.items():
        if n > 1:
            notes.append(f"obligor id '{base_id}' appears {n} times; later rows renamed "
                         f"'{base_id} #2'..'{base_id} #{n}' -- confirm they are distinct exposures")
    # review 2026-09-11 (C12): status values outside the defaulted list are Performing; list them.
    if "status" in cmap:
        seen_status = sorted({str(source_value(r, cmap, "status")).strip() for r in raw
                              if source_value(r, cmap, "status") not in (None, "")})
        other = [s for s in seen_status if s.lower() not in context["defaulted"] and s.lower() != "performing"]
        if other:
            notes.append(f"status values treated as Performing (not in defaulted_values): {other}")

    summed, stated, tie_status, tie_msg = tie_out_summary(out_rows, context, stops, notes)
    report = {
        "deal": context["meta"].get("deal"),
        "currency": context["meta"].get("currency"),
        "rows_in": len(raw),
        "rows_out": len(out_rows),
        "rows_skipped_nonobligor": skipped,
        "defaulted": sum(1 for r in out_rows if r["status"] == "Defaulted"),
        "sum_rona": summed,
        "stated_reference_notional": stated,
        "tie_out_status": tie_status,
        "tie_out": tie_msg,
        "stops": stops,
        "notes": notes,
    }
    return out_rows, report


def write_canonical(out_path, rows):
    with open(out_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=CANON_COLS)
        w.writeheader()
        for r in rows:
            w.writerow({k: ("" if r.get(k) is None else r.get(k)) for k in CANON_COLS})


def write_report(report_path, report):
    with open(report_path, "w", encoding="utf-8") as f:
        json.dump(report, f, indent=2)


def print_report(report):
    print(
        f"Normalized {report['rows_out']} obligor rows "
        f"({report['rows_skipped_nonobligor']} non-obligor skipped) -> {report['canonical']}"
    )
    print(f"  {report['tie_out']}")
    for nt in report["notes"]:
        print(f"  note: {nt}")
    if report["stops"]:
        print(f"\nSTOP ({len(report['stops'])}) -- escalate, do not fabricate:")
        for sgl in report["stops"]:
            print(f"  - {sgl}")


def parse_args(argv=None):
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--tape", required=True)
    ap.add_argument("--mapping", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--report", default=None)
    return ap.parse_args(argv)


def main(argv=None):
    args = parse_args(argv)
    mapping = load_mapping(args.mapping)
    raw, headers = read_rows(args.tape, mapping.get("source", {}))
    out_rows, report = normalize_rows(raw, headers, mapping)
    # review 2026-09-11 (C9): a STOPPED run must not leave a complete canonical file under the
    # requested name. The rows stay inspectable under a name no downstream step picks up.
    if report["stops"]:
        root, ext = os.path.splitext(args.out)
        canonical_path = f"{root}.STOPPED{ext or '.csv'}"
    else:
        canonical_path = args.out
    report["canonical"] = canonical_path

    write_canonical(canonical_path, out_rows)
    if args.report:
        write_report(args.report, report)
    print_report(report)
    if report["stops"]:
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
