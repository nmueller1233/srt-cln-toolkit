#!/usr/bin/env python3
"""Aggregate a reference registry into the CLN performance-curve inputs.

This is the registry->`cln.xlsx` bridge. The CLN file (`Performance Curves` tab,
one row per deal) takes RONA-weighted *distribution* inputs and computes CDR and
Severity itself via the `Parameters` reference tables. This module produces those
distribution inputs from a canonical registry, in the EXACT bucket order the file
expects, and -- because registries vary wildly in columns and granularity -- always
lands in one of two methodologically-consistent paths:

  PATH A (obligor):   the registry carries obligor-level PD and LGD/recovery
                      -> use them directly (RONA-weighted).
  PATH B (aggregate): it does not -> aggregate the pool into the rating / loan-type
                      / geography distributions and let the CLN engine derive
                      CDR + Severity.

Which path is *allowed* this quarter is governed by methodology consistency with the
prior period, NOT by what the new tape happens to contain (see `choose_path`).

The CDR/Severity formulas and the `Parameters` constants are mirrored here ONLY to
let the caller self-check the file's output and to anchor regression tests; the file
remains the source of truth. Verified against the live `cln.xlsx` named ranges:
  CDR      = MAX(minCDR, SUMPRODUCT(ratingDist, avgCDR[C]) * SUMPRODUCT(geoDist, geoMult)/100)
  Severity = 1 - SUMPRODUCT(loanDist, avgRecovery[C]) / 100

Limits stated by the 2026-09-11 review (each is filed in review/PROPOSALS.md, not changed here):
  - `aggregate()` does not filter on status. The obligor path excludes Defaulted rows by
    default; this path builds the distributions from every row it is given. Pass the
    performing rows if the deal's methodology excludes defaulted names from the forward
    CDR (the reference says it should) -- the caller decides, the function does not.
  - RONA that maps to no loan-type or region bucket stays in the denominator and in no
    bucket, so the emitted distributions can sum to less than 1. The file's formulas do
    not renormalise: unmapped loan type pushes Severity UP (toward 100 %), unmapped
    region pushes CDR DOWN (toward the floor). The shortfall is reported as a note and
    the sums are printed; it is not a STOP.
  - `write_cln_row` saves `cln.xlsx` with openpyxl. That preserves formulas but not
    add-in (CapIQ) connections or charts, so write into a COPY, never the live file.

Guidance, not deal rules: confirm thresholds and any deal-specific treatment with your reviewer.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys

# ---- Fixed bucket orders (must match cln.xlsx `Performance Curves` columns) -------
RATING_BUCKETS = ["IG", "BB", "B", "CCC", "NA"]            # D:H  (and J:N renormalized)
LOAN_TYPES = ["Leverage Loan", "Sn Unsecured", "Sn Secured",
              "Term Loans", "Revolver", "Trade Finance", "Property"]   # O:U
REGIONS = ["US", "Eur", "Middle East", "India", "EM/Asia"]            # V:Z

# ---- Parameters constants, mirrored from cln.xlsx (Parameters tab) ----------------
# Vintage: read from the live Parameters tab 2026-06-08. The only in-repo guard is the two
# golden rows in tests/srt/test_registry_granularity.py (0.2527 / 0.5229), which catch drift in
# THIS mirror, not a later change to the file. Re-read the tab when cln.xlsx changes.
# avgLoanCDR/avgBondCDR are percent default rates over [IG,BB,B,CCC,NA].
AVG_CDR = {
    "Loan": [0.00625, 0.38, 2.76, 7.897960500764167, 1.33],
    "Bond": [0.0, 0.525, 1.86, 7.675, 0.845],
    "Trade Finance": [None, None, None, None, None],   # all NA in the file
}
# avgRecoveryRate_* are percent recoveries over the 7 LOAN_TYPES.
AVG_RECOVERY = {
    "Loan": [61.2, 46.9, 61.2, 52.1, 80.6, 70.0, 80.0],
    "Bond": [None, 37.6, 54.8, None, None, None, None],
}
# geographicCDRMult = Parameters!AG4:AK4 (positional; see the label/multiplier note).
GEO_MULT = [1.0, 0.9733171619163131, 0.4730139478471801,
            0.5548817465130383, 0.6725288053365678]
MIN_CDR = 0.0025

# Heuristic keyword maps. Registries differ; these are the defaults a per-issuer
# mapping can extend. Order matters: the FIRST match wins, so more specific tokens
# (e.g. "leveraged") are listed before broader ones ("term").
RATING_TOKENS = {  # external scale -> bucket; longest-token-first match applied in code
    "AAA": "IG", "AA": "IG", "A": "IG", "BBB": "IG", "BAA": "IG",
    "BB": "BB", "BA": "BB", "B": "B",
    "CCC": "CCC", "CC": "CCC", "C": "CCC", "CAA": "CCC", "CA": "CCC", "D": "CCC",
}
LOAN_TYPE_TOKENS = [
    ("leverage", "Leverage Loan"), ("leveraged", "Leverage Loan"),
    ("trade finance", "Trade Finance"), ("trade-finance", "Trade Finance"),
    ("property", "Property"), ("real estate", "Property"), ("cre", "Property"),
    ("revolv", "Revolver"), ("rcf", "Revolver"),
    ("term", "Term Loans"), ("tlb", "Term Loans"), ("tla", "Term Loans"),
    ("senior secured", "Sn Secured"), ("sr secured", "Sn Secured"), ("1st lien", "Sn Secured"),
    ("senior unsecured", "Sn Unsecured"), ("sr unsecured", "Sn Unsecured"),
    ("unsecured", "Sn Unsecured"), ("secured", "Sn Secured"),
]
# Tokens that are also substrings of common words ("cre" in "credit") match as whole words only.
WHOLE_WORD_TOKENS = {"cre", "rcf", "tlb", "tla"}
REGION_TOKENS = [
    ("united states", "US"), ("usa", "US"), ("u.s", "US"), ("canada", "US"),
    ("north america", "US"), ("mexico", "US"),  # NB: file maps Latam-mult positionally
    ("united kingdom", "Eur"), ("uk", "Eur"), ("europe", "Eur"), ("eur", "Eur"),
    ("germany", "Eur"), ("france", "Eur"), ("spain", "Eur"), ("italy", "Eur"),
    ("netherlands", "Eur"), ("ireland", "Eur"), ("nordic", "Eur"), ("switzerland", "Eur"),
    ("india", "India"),
    ("middle east", "Middle East"), ("uae", "Middle East"), ("saudi", "Middle East"),
    ("qatar", "Middle East"), ("emirates", "Middle East"), ("gcc", "Middle East"),
    # everything else emerging/asian
    ("china", "EM/Asia"), ("asia", "EM/Asia"), ("emerging", "EM/Asia"),
    ("brazil", "EM/Asia"), ("latam", "EM/Asia"), ("africa", "EM/Asia"),
    ("indonesia", "EM/Asia"), ("singapore", "EM/Asia"), ("korea", "EM/Asia"),
]


def to_float(v):
    if v is None:
        return None
    s = str(v).strip().replace(",", "")
    pct = s.endswith("%")
    s = s.replace("%", "")
    if s == "":
        return None
    try:
        x = float(s)
    except ValueError:
        return None
    return x / 100.0 if pct else x


def bucket_rating(raw):
    """Map any external rating string to {IG,BB,B,CCC}. None/blank -> 'NA'."""
    if raw is None or str(raw).strip() == "":
        return "NA"
    s = str(raw).strip().upper()
    for token in sorted(RATING_TOKENS, key=len, reverse=True):
        if s.startswith(token):
            return RATING_TOKENS[token]
    return "NA"


def bucket_loan_type(raw):
    """First matching token wins. Short tokens (`cre`, `rcf`, `tlb`, `tla`) match only as whole
    words: review 2026-09-11 found `cre` matching the substring of "credit", so every
    "... Credit Facility" descriptor landed in Property (80 % recovery)."""
    if raw is None:
        return None
    s = str(raw).strip().lower()
    for token, bucket in LOAN_TYPE_TOKENS:
        if token in WHOLE_WORD_TOKENS:
            if re.search(r"(?<![a-z0-9])" + re.escape(token) + r"(?![a-z0-9])", s):
                return bucket
        elif token in s:
            return bucket
    return None


def bucket_region(raw):
    if raw is None:
        return None
    s = str(raw).strip().lower()
    for token, bucket in REGION_TOKENS:
        if token in s:
            return bucket
    return None


def pick_external_rating(row, rating_source_cols, fill_counts):
    """Per-obligor rating from the most-populated agency column, falling back across
    columns by descending fill count so coverage is maximized. Returns (raw, source)."""
    for col in sorted(rating_source_cols, key=lambda c: fill_counts.get(c, 0), reverse=True):
        val = row.get(col)
        if val is not None and str(val).strip() != "":
            return val, col
    return None, None


def column_fill_counts(rows, cols):
    return {c: sum(1 for r in rows if r.get(c) not in (None, "")) for c in cols}


def weighted_distribution(rows, weight_key, bucket_fn, buckets, value_key):
    """RONA-weighted proportion in each bucket. Unbucketed weight is returned separately."""
    totals = {b: 0.0 for b in buckets}
    unbucketed = 0.0
    total_w = 0.0
    for r in rows:
        w = to_float(r.get(weight_key))
        if w is None or w <= 0:
            continue
        total_w += w
        b = bucket_fn(r.get(value_key))
        if b in totals:
            totals[b] += w
        else:
            unbucketed += w
    if total_w <= 0:
        return {b: 0.0 for b in buckets}, 0.0, 0.0
    dist = {b: totals[b] / total_w for b in buckets}
    return dist, unbucketed / total_w, total_w


def severity_from_loan_dist(loan_dist, bond_loan_flag):
    rec = AVG_RECOVERY.get(bond_loan_flag)
    if rec is None:
        return None
    acc = 0.0
    for i, lt in enumerate(LOAN_TYPES):
        if rec[i] is None:
            if loan_dist.get(lt, 0.0) > 0:
                return None  # weight on a loan type with no recovery rate for this table -> file shows "Error"
            continue
        acc += loan_dist.get(lt, 0.0) * rec[i]
    return 1.0 - acc / 100.0


def cdr_from_dists(rating_dist, geo_dist, bond_loan_flag):
    cdr_rates = AVG_CDR.get(bond_loan_flag)
    if cdr_rates is None:
        return None
    rating_leg = 0.0
    for i, rb in enumerate(RATING_BUCKETS):
        if cdr_rates[i] is None:
            if rating_dist.get(rb, 0.0) > 0:
                return None
            continue
        rating_leg += rating_dist.get(rb, 0.0) * cdr_rates[i]
    geo_leg = sum(geo_dist.get(REGIONS[i], 0.0) * GEO_MULT[i] for i in range(len(REGIONS)))
    return max(MIN_CDR, rating_leg * geo_leg / 100.0)


def choose_path(prior_methodology, registry_has_obligor_pdlgd, new_deal):
    """Decide the methodologically-consistent path. Returns (path, stop, message).

    prior_methodology in {None, 'obligor', 'aggregate'}; None means no prior (new deal).
    """
    if new_deal or prior_methodology is None:
        if registry_has_obligor_pdlgd:
            return "obligor", False, (
                "New deal with obligor-level PD/LGD present -> obligor path. Record this "
                "as the calibrated methodology so future quarterly updates stay consistent.")
        return "aggregate", False, (
            "New deal without obligor-level PD/LGD -> aggregate to the CLN performance "
            "curve. Record aggregation as the calibrated methodology.")
    if prior_methodology == "obligor":
        if registry_has_obligor_pdlgd:
            return "obligor", False, "Prior methodology used obligor PD/LGD; new tape still provides them -> continue obligor path."
        return "obligor", True, (
            "STOP: prior methodology was obligor-level PD/LGD but the new tape lacks them. "
            "Do not silently switch to aggregation -- that injects a methodology-driven mark "
            "variance. Investigate source gap; human review.")
    if prior_methodology == "aggregate":
        if registry_has_obligor_pdlgd:
            return "aggregate", True, (
                "STOP: prior methodology was AGGREGATION but the new tape now carries obligor "
                "PD/LGD. Two consistent paths: (a) keep aggregating to avoid a methodology-driven "
                "variance in the submitted mark, or (b) switch to obligor-level and attribute the "
                "resulting change to the methodology shift, not to risk. Subject to human review; "
                "investigate divergence between the two before deciding.")
        return "aggregate", False, "Prior methodology used aggregation; new tape has no obligor PD/LGD -> continue aggregation."
    return "aggregate", True, f"Unknown prior methodology {prior_methodology!r} -> STOP and confirm with the prior-period evidence."


def internal_rating_check(used_internal_only, internal_mapping_present):
    """Flag generic internal ratings that cannot be tied to portfolio risk."""
    if used_internal_only and not internal_mapping_present:
        return True, (
            "STOP/FLAG: the only ratings available are INTERNAL grades with no mapping to an "
            "external/PD-calibrated scale. A generic internal rating cannot be bucketed to "
            "IG/BB/B/CCC without misstating risk. Help the analyst source an internal->external "
            "mapping or external ratings; do not guess a bucket.")
    return False, ""


def aggregate(rows, cfg):
    """Produce the CLN distribution inputs + self-checked CDR/Severity + flags."""
    weight_key = cfg.get("rona_col", "rona")
    bond_loan_flag = cfg.get("bond_loan_flag", "Loan")

    rating_cols = cfg.get("external_rating_cols", [])
    internal_cols = cfg.get("internal_rating_cols", [])
    internal_map = cfg.get("internal_to_external_map")  # dict or None
    fills = column_fill_counts(rows, rating_cols)

    used_internal_only = bool(internal_cols) and not rating_cols
    rated = []
    for r in rows:
        raw, src = pick_external_rating(r, rating_cols, fills)
        if raw is None and internal_cols and internal_map:
            for ic in internal_cols:
                iv = r.get(ic)
                if iv is not None and str(iv).strip() in internal_map:
                    raw = internal_map[str(iv).strip()]
                    break
        rr = dict(r)
        rr["_rating_resolved"] = raw
        rated.append(rr)

    flag_internal, internal_msg = internal_rating_check(used_internal_only, bool(internal_map))

    rating_dist, rating_unb, total_w = weighted_distribution(
        rated, weight_key, bucket_rating, RATING_BUCKETS, "_rating_resolved")
    loan_dist, loan_unb, _ = weighted_distribution(
        rows, weight_key, bucket_loan_type, LOAN_TYPES, cfg.get("loan_type_col", "loan_type"))
    geo_dist, geo_unb, _ = weighted_distribution(
        rows, weight_key, bucket_region, REGIONS, cfg.get("region_col", "country_domicile"))

    severity = severity_from_loan_dist(loan_dist, bond_loan_flag)
    cdr = cdr_from_dists(rating_dist, geo_dist, bond_loan_flag)

    stops, notes = [], []
    if flag_internal:
        stops.append(internal_msg)
    if loan_unb > 1e-9:
        notes.append(f"{loan_unb:.2%} of RONA could not be mapped to a loan type (loan-type distribution sums to "
                     f"{1 - loan_unb:.4f}, not 1). The file does not renormalise: this share is priced at 100% severity, "
                     f"so self_check_severity is biased UPWARD. Extend the mapping or record a judgment before writing.")
    if geo_unb > 1e-9:
        notes.append(f"{geo_unb:.2%} of RONA could not be mapped to a region (region distribution sums to "
                     f"{1 - geo_unb:.4f}, not 1). The file does not renormalise: this share carries a zero CDR multiplier, "
                     f"so self_check_cdr is biased DOWNWARD. Extend the mapping or record a judgment before writing.")
    if severity is None:
        stops.append(f"Severity not computable for '{bond_loan_flag}' table -- a loan-type bucket has no recovery rate (file would show 'Error').")
    notes.append("Geography label/multiplier note: cln.xlsx multiplies V:Z (US,Eur,Middle East,India,EM/Asia) "
                 "positionally by Parameters AG:AK (US,Eur,Latam,EM&FM,Asia). Labels differ; preserved as-is. Flag for human review.")

    return {
        "deal": cfg.get("deal"),
        "bond_loan_flag": bond_loan_flag,
        "include_na": "Yes",  # protocol: NA kept as its own bucket
        "rating_distribution": rating_dist,
        "loan_type_distribution": loan_dist,
        "region_distribution": geo_dist,
        "self_check_cdr": cdr,
        "self_check_severity": severity,
        "rona_total": total_w,
        "internal_rating_flag": flag_internal,
        "stops": stops,
        "notes": notes,
    }


def cln_row_inputs(result):
    """Map the aggregation result to the cln.xlsx `Performance Curves` input columns.
    Inputs only -- J:N (renorm), AB (CDR), AC (Severity) are FORMULAS; never write them."""
    rd, ld, gd = result["rating_distribution"], result["loan_type_distribution"], result["region_distribution"]
    return {
        "C": result["bond_loan_flag"],
        "D": rd["IG"], "E": rd["BB"], "F": rd["B"], "G": rd["CCC"], "H": rd["NA"],
        "I": result["include_na"],
        "O": ld["Leverage Loan"], "P": ld["Sn Unsecured"], "Q": ld["Sn Secured"],
        "R": ld["Term Loans"], "S": ld["Revolver"], "T": ld["Trade Finance"], "U": ld["Property"],
        "V": gd["US"], "W": gd["Eur"], "X": gd["Middle East"], "Y": gd["India"], "Z": gd["EM/Asia"],
    }


# Columns that are ALWAYS formulas in every Performance Curves row -- never write them
# as literals (renormalized rating dist + the CDR/Severity outputs).
FORMULA_COLS_NEVER_WRITE = {"J", "K", "L", "M", "N", "AB", "AC"}


def write_cln_row(cln_path, deal_name, row_inputs, cpr=0, notes=None,
                  sheet="Performance Curves", deal_col_letter="B", header_row=2):
    """Write a deal's distribution inputs into cln.xlsx, preserving every formula.

    Updates the deal's existing row if present, else APPENDS a new row below all
    content (never insert_rows -- openpyxl does not fix formula refs on insert). On
    append, the always-formula cells (J:N, AB, AC) are copied from the last data row
    and TRANSLATED to the new row so CDR/Severity compute; the distribution inputs
    (C, D:H, I, O:U, V:Z, AA) are written as values. Operate on a COPY, not the live
    file. Returns {row, appended, sheet}.
    """
    from openpyxl import load_workbook
    from openpyxl.utils import column_index_from_string, get_column_letter
    from openpyxl.formula.translate import Translator

    bad = [c for c in row_inputs if c in FORMULA_COLS_NEVER_WRITE]
    if bad:
        raise ValueError(f"refusing to write formula columns as values: {bad}")

    wb = load_workbook(cln_path)  # data_only=False -> keeps formulas
    if sheet not in wb.sheetnames:
        raise ValueError(f"sheet {sheet!r} not found in {wb.sheetnames}")
    ws = wb[sheet]
    deal_col = column_index_from_string(deal_col_letter)

    named_rows = [r for r in range(header_row + 1, ws.max_row + 1)
                  if ws.cell(r, deal_col).value not in (None, "")]
    last_content = header_row
    for r in range(header_row + 1, ws.max_row + 1):
        if any(ws.cell(r, c).value not in (None, "") for c in range(1, ws.max_column + 1)):
            last_content = r

    target = None
    for r in named_rows:
        if str(ws.cell(r, deal_col).value).strip().lower() == deal_name.strip().lower():
            target = r
            break
    appended = target is None
    if appended:
        if not named_rows:
            raise ValueError("no existing data row to use as a formula template")
        template = named_rows[-1]
        target = last_content + 1
        for c in range(1, ws.max_column + 1):
            tv = ws.cell(template, c).value
            if isinstance(tv, str) and tv.startswith("="):
                origin = f"{get_column_letter(c)}{template}"
                dest = f"{get_column_letter(c)}{target}"
                ws.cell(target, c).value = Translator(tv, origin=origin).translate_formula(dest)

    ws.cell(target, deal_col).value = deal_name
    for col, val in row_inputs.items():
        ws.cell(target, column_index_from_string(col)).value = val
    ws.cell(target, column_index_from_string("AA")).value = cpr
    if notes is not None:
        ws.cell(target, column_index_from_string("AD")).value = notes

    wb.save(cln_path)
    return {"row": target, "appended": appended, "sheet": sheet}


def read_registry(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--registry", required=True, help="canonical or raw registry CSV")
    ap.add_argument("--config", required=True, help="cln-mapping JSON (columns, flags, internal map)")
    ap.add_argument("--out", default=None, help="write the result JSON here")
    ap.add_argument("--cln", default=None, help="path to a COPY of cln.xlsx to write the row into")
    ap.add_argument("--cpr", type=float, default=0.0, help="CPR input for the row (default 0)")
    args = ap.parse_args(argv)
    cfg = json.load(open(args.config, encoding="utf-8"))
    rows = read_registry(args.registry)
    result = aggregate(rows, cfg)
    result["cln_row_inputs"] = cln_row_inputs(result)
    if args.cln:
        if result["stops"]:
            result["cln_write"] = "SKIPPED -- unresolved STOPs; not writing to cln.xlsx"
        else:
            result["cln_write"] = write_cln_row(
                args.cln, cfg.get("deal", "Unnamed Deal"), result["cln_row_inputs"],
                cpr=args.cpr, notes=cfg.get("write_note"))
    text = json.dumps(result, indent=2)
    if args.out:
        open(args.out, "w", encoding="utf-8").write(text)
    print(text)
    return 1 if result["stops"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
