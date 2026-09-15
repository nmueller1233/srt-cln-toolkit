"""CLN writer contract test (skippable, like the model-map validator).

Runs ONLY on a machine that has openpyxl and the live cln.xlsx (path from
corpus.map.json: corpus_root + model_skeleton.cln_performance_curve). Everywhere else
(CI, plugin-eval, a teammate without the corpus) it SKIPs.

It copies cln.xlsx to a temp file, writes a synthetic deal row via
aggregate_to_cln.write_cln_row, reloads, and asserts the contract:
  - a NEW row was appended below existing content;
  - the distribution INPUTS landed as values;
  - the always-formula cells (J:N renorm, AB CDR, AC Severity) are present as formulas
    TRANSLATED to the new row -- never written as literals;
  - the existing rows (e.g. the first data row) are untouched.
The live file is never modified -- only the temp copy.
"""
import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

PLUGIN_ROOT = Path(__file__).resolve().parents[2]
# registry scripts live under scripts/srt-registry-inputs/ in the plugin.
for _cand in (PLUGIN_ROOT / "scripts" / "srt-registry-inputs", PLUGIN_ROOT / "scripts"):
    if (_cand / "aggregate_to_cln.py").exists():
        sys.path.insert(0, str(_cand))
        break

try:
    import openpyxl  # noqa: F401
    from openpyxl.utils import column_index_from_string
    import aggregate_to_cln as agg
    HAVE_OPENPYXL = True
except Exception:
    HAVE_OPENPYXL = False


def _cln_path():
    cmap = PLUGIN_ROOT / "assets" / "srt" / "corpus.map.json"
    if not cmap.exists():
        return None
    cfg = json.loads(cmap.read_text(encoding="utf-8"))
    root = Path(cfg.get("corpus_root", ""))
    rel = cfg.get("model_skeleton", {}).get("cln_performance_curve")
    if not rel:
        return None
    return root / rel


CLN = _cln_path() if HAVE_OPENPYXL else None
WHY = "needs openpyxl + the live cln.xlsx from corpus.map.json"


@unittest.skipUnless(HAVE_OPENPYXL and CLN and CLN.exists(), WHY)
class ClnWriteContractTests(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp(prefix="cln_write_")
        self.tmp = Path(self.tmpdir) / "cln-copy.xlsx"
        shutil.copy2(CLN, self.tmp)
        self.deal = "ZZ Synthetic Contract-Test Deal"
        # 100% IG / 100% Term Loans / 100% US, Loan flag
        self.inputs = {"C": "Loan", "D": 1.0, "E": 0.0, "F": 0.0, "G": 0.0, "H": 0.0,
                       "I": "Yes", "O": 0.0, "P": 0.0, "Q": 0.0, "R": 1.0, "S": 0.0,
                       "T": 0.0, "U": 0.0, "V": 1.0, "W": 0.0, "X": 0.0, "Y": 0.0, "Z": 0.0}

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def _reload(self):
        return openpyxl.load_workbook(self.tmp)["Performance Curves"]

    def test_appends_row_with_values_and_translated_formulas(self):
        before = self._reload()
        existing_d = before["D3"].value  # capture an existing row to prove it is untouched

        res = agg.write_cln_row(str(self.tmp), self.deal, self.inputs, cpr=0.0,
                                notes="contract test")
        self.assertTrue(res["appended"])
        row = res["row"]

        ws = self._reload()
        C = column_index_from_string
        # deal name + inputs landed as values
        self.assertEqual(ws.cell(row, C("B")).value, self.deal)
        self.assertEqual(ws.cell(row, C("C")).value, "Loan")
        self.assertEqual(ws.cell(row, C("I")).value, "Yes")
        self.assertAlmostEqual(ws.cell(row, C("D")).value, 1.0)  # IG
        self.assertAlmostEqual(ws.cell(row, C("R")).value, 1.0)  # Term Loans
        self.assertAlmostEqual(ws.cell(row, C("V")).value, 1.0)  # US

        # always-formula cells present, TRANSLATED to this row, not literals
        for col in ("J", "K", "L", "M", "N", "AB", "AC"):
            v = ws.cell(row, C(col)).value
            self.assertIsInstance(v, str, f"{col}{row} should be a formula")
            self.assertTrue(v.startswith("="), f"{col}{row} should start with '='")
        self.assertIn(f"O{row}:U{row}", ws.cell(row, C("AC")).value)   # severity refs this row
        self.assertIn(f"J{row}:N{row}", ws.cell(row, C("AB")).value)   # CDR refs this row

        # existing row untouched
        self.assertEqual(ws["D3"].value, existing_d)

    def test_refuses_to_write_formula_columns_as_values(self):
        with self.assertRaises(ValueError):
            agg.write_cln_row(str(self.tmp), self.deal, {"AC": 0.25})


if __name__ == "__main__":
    unittest.main()
